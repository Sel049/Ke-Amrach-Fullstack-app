#!/usr/bin/env python3
"""
PowerPoint Presentation Generator for Ethio-Farmers-Shop Project
Converts markdown content to a professional PowerPoint presentation
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import re

def create_presentation():
    """Create the PowerPoint presentation"""
    
    # Create presentation object
    prs = Presentation()
    
    # Set slide dimensions (16:9 aspect ratio)
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    
    # Read the markdown content
    with open('Ethio-Farmers-Shop-Presentation.md', 'r', encoding='utf-8') as file:
        content = file.read()
    
    # Split content into slides
    slides_content = content.split('---')
    
    for slide_content in slides_content:
        if not slide_content.strip():
            continue
            
        # Extract slide title and content
        lines = slide_content.strip().split('\n')
        if len(lines) < 2:
            continue
            
        # Skip the markdown header line
        if lines[0].startswith('## Slide'):
            lines = lines[1:]
        
        # Find title (first line that's not empty)
        title = ""
        content_start = 0
        for i, line in enumerate(lines):
            if line.strip() and not line.startswith('**') and not line.startswith('-') and not line.startswith('📊') and not line.startswith('🌾') and not line.startswith('👨') and not line.startswith('🛒') and not line.startswith('🔗') and not line.startswith('🔐') and not line.startswith('🎨') and not line.startswith('📱') and not line.startswith('🏦') and not line.startswith('🛡️') and not line.startswith('📈') and not line.startswith('👨‍💼') and not line.startswith('🔄') and not line.startswith('⚡') and not line.startswith('☁️') and not line.startswith('🤖') and not line.startswith('🏛️') and not line.startswith('✅') and not line.startswith('🔧') and not line.startswith('🧪') and not line.startswith('📚') and not line.startswith('🇪🇹') and not line.startswith('🎯') and not line.startswith('🔒') and not line.startswith('💻') and not line.startswith('📊') and not line.startswith('🚀') and not line.startswith('📧') and not line.startswith('🤝'):
                title = line.strip()
                content_start = i + 1
                break
        
        if not title:
            title = "Slide"
        
        # Create slide
        slide_layout = prs.slide_layouts[1]  # Title and Content layout
        slide = prs.slides.add_slide(slide_layout)
        
        # Set title
        title_shape = slide.shapes.title
        title_shape.text = title
        title_shape.text_frame.paragraphs[0].font.size = Pt(28)
        title_shape.text_frame.paragraphs[0].font.bold = True
        title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 100, 0)  # Dark green
        
        # Set content
        content_shape = slide.placeholders[1]
        text_frame = content_shape.text_frame
        text_frame.clear()
        
        # Process content lines
        current_paragraph = None
        for line in lines[content_start:]:
            line = line.strip()
            if not line:
                continue
                
            # Handle different content types
            if line.startswith('**') and line.endswith('**'):
                # Bold text
                if current_paragraph is None:
                    current_paragraph = text_frame.add_paragraph()
                current_paragraph.text = line[2:-2]
                current_paragraph.font.bold = True
                current_paragraph.font.size = Pt(18)
                current_paragraph = None
            elif line.startswith('- '):
                # Bullet point
                if current_paragraph is None:
                    current_paragraph = text_frame.add_paragraph()
                current_paragraph.text = line[2:]
                current_paragraph.font.size = Pt(14)
                current_paragraph.level = 0
                current_paragraph = None
            elif line.startswith('  - '):
                # Sub bullet point
                if current_paragraph is None:
                    current_paragraph = text_frame.add_paragraph()
                current_paragraph.text = line[4:]
                current_paragraph.font.size = Pt(12)
                current_paragraph.level = 1
                current_paragraph = None
            elif line.startswith('🌾') or line.startswith('🛒') or line.startswith('🔗') or line.startswith('🔐') or line.startswith('🎨') or line.startswith('📱') or line.startswith('🏦') or line.startswith('🛡️') or line.startswith('📈') or line.startswith('👨') or line.startswith('🔄') or line.startswith('⚡') or line.startswith('☁️') or line.startswith('🤖') or line.startswith('🏛️') or line.startswith('✅') or line.startswith('🔧') or line.startswith('🧪') or line.startswith('📚') or line.startswith('🇪🇹') or line.startswith('🎯') or line.startswith('🔒') or line.startswith('💻') or line.startswith('📊') or line.startswith('🚀') or line.startswith('📧') or line.startswith('🤝'):
                # Emoji section header
                if current_paragraph is None:
                    current_paragraph = text_frame.add_paragraph()
                current_paragraph.text = line
                current_paragraph.font.bold = True
                current_paragraph.font.size = Pt(16)
                current_paragraph.font.color.rgb = RGBColor(0, 100, 0)  # Dark green
                current_paragraph = None
            else:
                # Regular text
                if current_paragraph is None:
                    current_paragraph = text_frame.add_paragraph()
                current_paragraph.text = line
                current_paragraph.font.size = Pt(14)
                current_paragraph = None
    
    # Save the presentation
    prs.save('Ethio-Farmers-Shop-Presentation.pptx')
    print("✅ PowerPoint presentation created successfully!")
    print("📁 File saved as: Ethio-Farmers-Shop-Presentation.pptx")

if __name__ == "__main__":
    try:
        create_presentation()
    except ImportError:
        print("❌ Error: python-pptx library not found.")
        print("📦 Please install it using: pip install python-pptx")
    except FileNotFoundError:
        print("❌ Error: Ethio-Farmers-Shop-Presentation.md file not found.")
        print("📁 Please make sure the markdown file exists in the current directory.")
    except Exception as e:
        print(f"❌ Error creating presentation: {e}")
